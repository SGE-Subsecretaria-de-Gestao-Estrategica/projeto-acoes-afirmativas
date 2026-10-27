"""Gabarito manual e F1 da extração.

1) gera_planilha_revisao(): uma linha por edital x campo, com o valor do modelo, a evidência
   e as colunas em branco para a revisão manual (modelo_correto S/N e valor_real).
2) calcula_f1(): lê a planilha revisada e calcula precisão, revocação e F1 por campo.
"""
from pathlib import Path
from urllib.parse import quote

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from config import settings
from dicionario_cotas import menciona

AMOSTRA = settings.LPG_INTERIM / "amostra_gabarito.csv"
REVISAO = settings.FINAL_DATA_PATH / "gabarito_revisao_lpg.xlsx"
F1_SAIDA = settings.FINAL_DATA_PATH / "f1_lpg.xlsx"

# (coluna na planilha final, campo da evidência, o que conferir no edital)
CAMPOS = [
    ("vagas_totais", "vagas_totais", "Total de vagas/projetos/prêmios do edital"),
    ("valor_total", "valor_total", "Valor total do edital, em reais"),
    ("possui_categorias", "categorias", "O edital é dividido em categorias/módulos/linhas com vagas próprias?"),
    ("numero_vagas_categoria_total", "categorias", "Soma das vagas de todas as categorias do edital"),
    ("perc_cotas_negras", "percentual_negras", "Percentual de cota para pessoas negras, em decimal (0.2)"),
    ("vagas_cotas_negras", "vagas_negras_total", "Vagas reservadas a pessoas negras no edital (soma das categorias, ou percentual × vagas com o arredondamento da IN)"),
    ("perc_cotas_indigenas", "percentual_indigenas", "Percentual de cota para pessoas indígenas, em decimal (0.1)"),
    ("vagas_cotas_indigenas", "vagas_indigenas_total", "Vagas reservadas a pessoas indígenas no edital (soma das categorias, ou percentual × vagas com o arredondamento da IN)"),
    ("edital_especifico_grupo", "edital_especifico_grupo", "O edital inteiro é para um grupo? Qual?"),
    ("outras_cotas", "outras_cotas", "Cotas (reserva de vagas) para outros grupos além de negros e indígenas, com o percentual. Bônus na nota não conta"),
]
CAMPOS_SIM_NAO = {"possui_categorias"}

# Link do PDF relativo à planilha. Funciona no SharePoint/OneDrive se a estrutura for
#   <pasta>/validação/gabarito_revisao_lpg.xlsx  e  <pasta>/editais/<Capitais|Estados>/<ente>/<arquivo>.pdf
# e também localmente (data/final -> data/EDITAIS), para os editais da base.
LINK_RELATIVO = "../editais/{pasta_tipo}/{ente}/{arquivo}"
LINK_RELATIVO_NOVOS = "../editais/NOVOS/{caminho}"  # data/NOVOS/<subpasta>/<arquivo> -> editais/NOVOS/...


def link_pdf(row: pd.Series) -> str:
    if row.get("origem") == "novos":
        caminho = Path(row["pdf_path"]).resolve().relative_to(settings.LPG_NOVOS_ROOT.resolve()).as_posix()
        rel = LINK_RELATIVO_NOVOS.format(caminho=caminho)
    else:
        rel = LINK_RELATIVO.format(
            pasta_tipo="Capitais" if row["tipo_ente"] == "CAPITAL" else "Estados", ente=row["ente"], arquivo=row["arquivo"],
        )
    if settings.SHAREPOINT_PASTA_URL:  # link absoluto: funciona para quem tem acesso à pasta, de onde quer que abra
        return settings.SHAREPOINT_PASTA_URL.rstrip("/") + "/" + quote(rel.removeprefix("../"))
    return rel
NULOS = {"", "nan", "none", "null", "sem_categorias", "nao_aplicavel"}  # 0 é valor, não ausência


def _vazio(v, campo: str = "") -> bool:
    if campo in CAMPOS_SIM_NAO:  # positivo = "sim"
        return str(v).strip().lower() != "sim"
    return v is None or (isinstance(v, float) and pd.isna(v)) or str(v).strip().lower() in NULOS


def _exibe(valor, campo: str):
    if pd.isna(valor):
        return ""
    if campo in CAMPOS_SIM_NAO:
        return "sim" if bool(valor) else "não"
    return valor


def _evidencia(evid: pd.DataFrame, edital_id: str, campo_ev: str | None) -> tuple[str, str]:
    if campo_ev is None:
        return "", ""
    ev = evid[evid["id"] == edital_id]
    if campo_ev in ("pcd", "lgbtqia", "quilombolas"):  # evidências de outras cotas: filtra pelo grupo
        ev = ev[(ev["campo"] == campo_ev) | ((ev["campo"] == "outras_cotas") & ev["trecho"].map(lambda t: menciona(str(t), campo_ev)))]
    else:
        ev = ev[ev["campo"] == campo_ev]
    if ev.empty:
        return "", ""
    r = ev.iloc[0]
    return str(r["trecho"]), "" if pd.isna(r["pagina"]) else str(int(r["pagina"]))


def gera_planilha_revisao(sobrescrever: bool = False) -> Path:
    if REVISAO.exists() and not sobrescrever:
        raise FileExistsError(f"{REVISAO} já existe (pode conter revisão manual). Use sobrescrever=True.")
    x = pd.read_excel(settings.LPG_OUTPUT, sheet_name=None)
    editais, evid = x["editais"], x["evidencias"]
    amostra = pd.read_csv(AMOSTRA, sep=";", dtype=str)
    sel = editais[editais["id"].isin(amostra["id"])].sort_values(["tipo_ente", "ente", "arquivo_pdf"])

    linhas = []
    for n, (_, e) in enumerate(sel.iterrows(), start=1):
        for col, campo_ev, pergunta in CAMPOS:
            trecho, pagina = _evidencia(evid, e["id"], campo_ev)
            valor = e.get(col)
            linhas.append({
                "n": n, "id": e["id"], "tipo_ente": e["tipo_ente"], "ente": e["ente"], "arquivo_pdf": e["arquivo_pdf"],
                "campo": col, "o_que_conferir": pergunta,
                "valor_modelo": _exibe(valor, col),
                "evidencia_trecho": trecho, "evidencia_pagina": pagina,
                "modelo_correto": "", "valor_real": "", "comentario": "",
            })
    df = pd.DataFrame(linhas)

    instrucoes = pd.DataFrame({"Como revisar": [
        "1. Para cada edital (coluna n), abra o PDF pelo link da coluna arquivo_pdf (link relativo: funciona com a planilha em 'validação' e os PDFs em 'editais/<Capitais|Estados>/<ente>/' ou 'editais/NOVOS/').",
        "2. Em cada linha, compare valor_modelo com o edital. A evidência e a página ajudam a localizar.",
        "3. modelo_correto: S se o valor do modelo está certo (inclusive vazio, quando o edital realmente não traz a informação); N se está errado.",
        "4. valor_real: preencha só quando modelo_correto = N, com o valor correto. Se a informação não existe no edital (o modelo inventou), escreva 'nenhum' - senão o erro conta como FP e FN. Nos campos sim/não, escreva 'sim' ou 'não'.",
        "5. Percentuais em decimal (20% -> 0.2). Valores em reais como número (1234567.89).",
        "6. Cota = reserva de vagas. Pontuação extra/bônus para um grupo NÃO é cota.",
        "7. outras_cotas: se errado, liste em valor_real os grupos com cota e o percentual (ex.: PcD 0.05; quilombolas 0.05), ou 'nenhum'.",
        f"Amostra: {sel['id'].nunique()} editais sorteados (estratificados por ente), {len(df)} linhas.",
    ]})

    REVISAO.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(REVISAO, engine="openpyxl") as xw:
        instrucoes.to_excel(xw, sheet_name="instrucoes", index=False)
        df.to_excel(xw, sheet_name="revisao", index=False)
    _formata(df, amostra.set_index("id").apply(link_pdf, axis=1))
    print(f"📝 {REVISAO} ({sel['id'].nunique()} editais, {len(df)} linhas)")
    return REVISAO


def _formata(df: pd.DataFrame, links: pd.Series) -> None:
    wb = load_workbook(REVISAO)
    ws = wb["revisao"]
    cols = {c: i + 1 for i, c in enumerate(df.columns)}
    larguras = {"n": 5, "id": 10, "tipo_ente": 10, "ente": 16, "arquivo_pdf": 30, "campo": 22, "o_que_conferir": 45,
                "valor_modelo": 35, "evidencia_trecho": 60, "evidencia_pagina": 9, "modelo_correto": 10,
                "valor_real": 25, "comentario": 35}
    for c, w in larguras.items():
        ws.column_dimensions[ws.cell(1, cols[c]).column_letter].width = w
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "H2"

    cinza, amarelo = PatternFill("solid", fgColor="EEEEEE"), PatternFill("solid", fgColor="FFF2CC")
    for r in range(2, ws.max_row + 1):
        par = int(ws.cell(r, cols["n"]).value) % 2 == 0
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(r, c)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if par:
                cell.fill = cinza
        for c in ("modelo_correto", "valor_real", "comentario"):
            ws.cell(r, cols[c]).fill = amarelo
        link = ws.cell(r, cols["arquivo_pdf"])
        link.hyperlink = links[ws.cell(r, cols["id"]).value]
        link.font = Font(color="0563C1", underline="single")

    dv = DataValidation(type="list", formula1='"S,N"', allow_blank=True)
    ws.add_data_validation(dv)
    letra = ws.cell(1, cols["modelo_correto"]).column_letter
    dv.add(f"{letra}2:{letra}{ws.max_row}")
    ws.auto_filter.ref = ws.dimensions
    wb["instrucoes"].column_dimensions["A"].width = 140
    wb.save(REVISAO)


def calcula_f1() -> pd.DataFrame:
    """Só entram as linhas com modelo_correto = S ou N (editais fora do escopo ficam em branco).
    TP: modelo trouxe valor e está certo | TN: vazio e de fato não existe |
    FN: deixou vazio um valor que existe |
    FP + FN: trouxe valor errado (errou e perdeu o valor verdadeiro) |
    só FP: trouxe valor, mas valor_real = 'nenhum' (a informação não existe no edital)."""
    df = pd.read_excel(REVISAO, sheet_name="revisao", dtype=str).fillna("")
    df["modelo_correto"] = df["modelo_correto"].str.strip().str.upper()
    revisado = df[df["modelo_correto"].isin(["S", "N"])]
    pendentes = (df["modelo_correto"] == "").sum()

    def classifica(r):
        tem_modelo = not _vazio(r["valor_modelo"], r["campo"])
        if r["modelo_correto"] == "S":
            return ("TP",) if tem_modelo else ("TN",)
        if tem_modelo:
            return ("FP",) if r["valor_real"].strip().lower() == "nenhum" else ("FP", "FN")
        return ("FN",)

    contagem = []
    for _, r in revisado.iterrows():
        for t in classifica(r):
            contagem.append({"campo": r["campo"], "tipo": t})
    c = pd.DataFrame(contagem).pivot_table(index="campo", columns="tipo", aggfunc="size", fill_value=0)
    for t in ("TP", "FP", "FN", "TN"):
        if t not in c:
            c[t] = 0
    total = c[["TP", "FP", "FN", "TN"]].sum().rename("TOTAL")
    c = pd.concat([c, total.to_frame().T])
    c["precisao"] = c["TP"] / (c["TP"] + c["FP"])
    c["revocacao"] = c["TP"] / (c["TP"] + c["FN"])
    c["f1"] = 2 * c["precisao"] * c["revocacao"] / (c["precisao"] + c["revocacao"])
    acerto = revisado.groupby("campo")["modelo_correto"].apply(lambda s: (s == "S").mean())
    c["acuracia"] = acerto.reindex(c.index)
    c.loc["TOTAL", "acuracia"] = (revisado["modelo_correto"] == "S").mean()
    c = c[["TP", "FP", "FN", "TN", "precisao", "revocacao", "f1", "acuracia"]].round(3)

    fora = df.groupby("n")["modelo_correto"].apply(lambda m: (m == "").all())
    editais_fora = df[df["n"].isin(fora[fora].index)].drop_duplicates("n")[["n", "ente", "arquivo_pdf", "comentario"]]
    erros = revisado[revisado["modelo_correto"] == "N"][["n", "ente", "arquivo_pdf", "campo", "valor_modelo", "valor_real", "comentario"]]
    with pd.ExcelWriter(F1_SAIDA, engine="openpyxl") as xw:
        c.to_excel(xw, sheet_name="f1")
        erros.to_excel(xw, sheet_name="erros", index=False)
        editais_fora.to_excel(xw, sheet_name="editais_fora", index=False)
    print(f"Editais avaliados: {revisado['n'].nunique()} | fora da avaliação: {len(editais_fora)} ({', '.join(editais_fora['ente'] + ' ' + editais_fora['arquivo_pdf'])})")
    print(f"Linhas revisadas: {len(revisado)} de {len(df)} ({pendentes} pendentes)")
    print(c.to_string())
    print(f"\n📊 {F1_SAIDA}")
    return c
