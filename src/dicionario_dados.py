"""Dicionário de dados da planilha final (data/final/lpg_editais.xlsx).

Gera data/final/dicionario_dados_lpg_editais.xlsx e avisa se alguma coluna da planilha
não estiver documentada aqui.
"""
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font

from config import settings

SAIDA = settings.FINAL_DATA_PATH / "dicionario_dados_lpg_editais.xlsx"

# Fonte: metadado (arquivo/pasta) | modelo (extraído pelo LLM do texto do edital) |
#        calculado (regras em Python a partir do que o modelo extraiu) | regex (dicionário de termos)
IDENT = [
    ("id", "texto", "Identificador do edital (hash do caminho do PDF). Liga as linhas de todas as abas.", "16 caracteres hexadecimais", "metadado"),
    ("tipo_ente", "texto", "Tipo de ente federativo que publicou o edital.", "CAPITAL / ESTADO", "metadado"),
    ("ente", "texto", "Nome do estado, do DF ou da capital (nome da pasta de origem).", "", "metadado"),
    ("arquivo_pdf", "texto", "Nome do arquivo PDF do edital.", "ex.: 2023-01_EST-BAHIA_LPG.pdf", "metadado"),
]

DICIONARIO = {
    "editais": IDENT + [
        ("cod_ibge", "texto", "Código IBGE do ente: 2 dígitos para UF, 7 para município.", "", "metadado"),
        ("origem", "texto", "Lote de coleta do edital.", "base (data/EDITAIS) / novos (data/NOVOS)", "metadado"),
        ("modo_leitura", "texto", "Como o PDF foi lido pelo modelo.", "texto (convertido para Markdown) / pdf (escaneado, lido como imagem)", "metadado"),
        ("paginas", "inteiro", "Número de páginas do PDF.", "", "metadado"),
        ("status_extracao", "texto", "Resultado da extração pelo modelo. Só 'ok' tem dados.", "ok / truncado (resposta passou do limite de tokens) / invalido / recusa / nao_extraido", "calculado"),
        ("erro_extracao", "texto", "Mensagem de erro quando status_extracao não é 'ok'.", "", "calculado"),
        ("tipo_documento", "texto", "Tipo de documento. Retificação que republica o edital completo conta como edital.", "edital / retificacao / resultado / anexo / outro", "modelo"),
        ("objeto", "texto", "Objeto do edital, resumido em até 2 frases.", "", "modelo"),
        ("modalidade", "texto", "Modalidade de fomento.", "fomento_execucao / premio / bolsa / subsidio / outro", "modelo"),
        ("area_cultural", "texto", "Área cultural do edital, quando informada.", "ex.: audiovisual, demais áreas", "modelo"),
        ("valor_total", "decimal", "Valor total do edital.", "reais (R$)", "modelo"),
        ("vagas_totais", "inteiro", "Total de vagas (projetos, propostas, prêmios ou bolsas) a selecionar.", "", "modelo"),
        ("possui_categorias", "booleano", "O edital é dividido em categorias, módulos, linhas ou faixas com vagas próprias.", "VERDADEIRO / FALSO", "modelo"),
        ("n_categorias", "inteiro", "Número de categorias listadas na aba 'categorias'.", "", "calculado"),
        ("numero_vagas_categoria_total", "inteiro", "Soma das vagas de todas as categorias do edital (aba 'categorias'). Vazio se o edital não tem categorias com vagas. Pode diferir de vagas_totais quando o edital é inconsistente.", "", "calculado"),
        ("cotas_por_categoria", "texto", "Se o edital aplica as cotas em cada categoria (IN 5, art. 5º, §3º).", "sim / nao / nao_informado / sem_categorias", "modelo"),
        ("edital_especifico_grupo", "texto", "Grupo ao qual o edital inteiro é destinado, quando é um edital específico.", "ex.: povos indígenas; vazio se não for específico", "modelo"),
        ("perc_cotas_negras", "decimal", "Percentual de vagas reservadas a pessoas negras (pretas e pardas).", "fração: 0.2 = 20%", "modelo"),
        ("vagas_cotas_negras", "inteiro", "Vagas reservadas a pessoas negras: soma informada por categoria ou, se o edital só traz o percentual, percentual × vagas com o arredondamento do §2º (por categoria, quando houver).", "", "calculado"),
        ("vagas_min_negras", "inteiro", "Mínimo exigido pela IN 5: 20% das vagas totais, com o arredondamento do §2º.", "", "calculado"),
        ("status_negras", "texto", "Cumprimento da cota para pessoas negras.", "conforme / nao_conforme / indeterminado / especifico", "calculado"),
        ("motivo_negras", "texto", "Justificativa de status_negras.", "", "calculado"),
        ("perc_cotas_indigenas", "decimal", "Percentual de vagas reservadas a pessoas indígenas.", "fração: 0.1 = 10%", "modelo"),
        ("vagas_cotas_indigenas", "inteiro", "Vagas reservadas a pessoas indígenas (mesma regra de vagas_cotas_negras).", "", "calculado"),
        ("vagas_min_indigenas", "inteiro", "Mínimo exigido pela IN 5: 10% das vagas totais, com o arredondamento do §2º.", "", "calculado"),
        ("status_indigenas", "texto", "Cumprimento da cota para pessoas indígenas.", "conforme / nao_conforme / indeterminado / especifico", "calculado"),
        ("motivo_indigenas", "texto", "Justificativa de status_indigenas.", "", "calculado"),
        ("status_in5", "texto", "Classificação final do edital frente à IN MinC nº 5/2023, art. 5º. nao_conforme se algum grupo é nao_conforme; indeterminado se falta informação; especifico se o edital é dirigido a um grupo; nao_aplicavel se o documento não é edital.", "conforme / nao_conforme / indeterminado / especifico / nao_aplicavel", "calculado"),
        ("motivo_in5", "texto", "Justificativa de status_in5 (motivos dos grupos não conformes).", "", "calculado"),
        ("valor_por_vaga", "decimal", "Valor médio por vaga: valor_total / vagas_totais.", "reais (R$)", "calculado"),
        ("valor_cotas_negras", "decimal", "Valor estimado destinado às cotas para pessoas negras: vagas_cotas_negras × valor_por_vaga.", "reais (R$)", "calculado"),
        ("valor_cotas_indigenas", "decimal", "Valor estimado destinado às cotas para pessoas indígenas: vagas_cotas_indigenas × valor_por_vaga.", "reais (R$)", "calculado"),
        ("outras_cotas", "texto", "Resumo das reservas de vagas para outros grupos, com o percentual entre parênteses (detalhe na aba 'outras_cotas').", "ex.: pessoas com deficiência (0.05)", "modelo"),
        ("tem_cota_pcd", "booleano", "Há reserva de vagas para pessoas com deficiência.", "VERDADEIRO / FALSO", "calculado"),
        ("perc_cotas_pcd", "decimal", "Percentual reservado a pessoas com deficiência.", "fração: 0.05 = 5%", "modelo"),
        ("vagas_cotas_pcd", "inteiro", "Vagas reservadas a pessoas com deficiência, quando informadas.", "", "modelo"),
        ("tem_cota_lgbtqia", "booleano", "Há reserva de vagas para pessoas LGBTQIA+ (inclui pessoas trans e travestis).", "VERDADEIRO / FALSO", "calculado"),
        ("perc_cotas_lgbtqia", "decimal", "Percentual reservado a pessoas LGBTQIA+.", "fração", "modelo"),
        ("vagas_cotas_lgbtqia", "inteiro", "Vagas reservadas a pessoas LGBTQIA+, quando informadas.", "", "modelo"),
        ("tem_cota_quilombolas", "booleano", "Há reserva de vagas para quilombolas.", "VERDADEIRO / FALSO", "calculado"),
        ("perc_cotas_quilombolas", "decimal", "Percentual reservado a quilombolas.", "fração", "modelo"),
        ("vagas_cotas_quilombolas", "inteiro", "Vagas reservadas a quilombolas, quando informadas.", "", "modelo"),
        ("pontuacao_bonus", "texto", "Grupos que recebem pontuação extra ou bônus na avaliação. Não é reserva de vagas.", "lista separada por ';'", "modelo"),
        ("outros_editais_no_pdf", "texto", "Outros editais contidos no mesmo PDF (ex.: Diário Oficial) e não extraídos.", "lista separada por ';'", "modelo"),
        ("observacoes", "texto", "Ambiguidades do edital relevantes para a análise de cotas, apontadas pelo modelo.", "", "modelo"),
        ("revisar_grupos", "texto", "Grupos que o dicionário de termos encontrou perto de 'cota/reserva' no texto, mas para os quais o modelo não registrou cota nem bônus. Indica revisão manual; pode ser falso positivo. Só para editais lidos como texto.", "negras, indigenas, pcd, lgbtqia, quilombolas", "regex"),
        ("trecho_regex_negras", "texto", "Trecho do edital que gerou o alerta de revisão para pessoas negras.", "", "regex"),
        ("trecho_regex_indigenas", "texto", "Trecho do edital que gerou o alerta de revisão para pessoas indígenas.", "", "regex"),
        ("trecho_regex_pcd", "texto", "Trecho do edital que gerou o alerta de revisão para PcD.", "", "regex"),
        ("trecho_regex_lgbtqia", "texto", "Trecho do edital que gerou o alerta de revisão para LGBTQIA+.", "", "regex"),
        ("trecho_regex_quilombolas", "texto", "Trecho do edital que gerou o alerta de revisão para quilombolas.", "", "regex"),
    ],
    "categorias": IDENT + [
        ("nome", "texto", "Nome da categoria, módulo, linha ou faixa, como no edital.", "", "modelo"),
        ("vagas", "inteiro", "Vagas da categoria.", "", "modelo"),
        ("valor_por_vaga", "decimal", "Valor unitário por projeto ou prêmio na categoria.", "reais (R$)", "modelo"),
        ("valor_categoria", "decimal", "Valor total da categoria.", "reais (R$)", "modelo"),
        ("vagas_negras", "inteiro", "Vagas reservadas a pessoas negras na categoria, quando o edital informa o número.", "", "modelo"),
        ("vagas_indigenas", "inteiro", "Vagas reservadas a pessoas indígenas na categoria, quando o edital informa o número.", "", "modelo"),
        ("vagas_pcd", "inteiro", "Vagas reservadas a pessoas com deficiência na categoria, quando informadas.", "", "modelo"),
        ("vagas_lgbtqia", "inteiro", "Vagas reservadas a pessoas LGBTQIA+ na categoria, quando informadas.", "", "modelo"),
        ("vagas_quilombolas", "inteiro", "Vagas reservadas a quilombolas na categoria, quando informadas.", "", "modelo"),
    ],
    "outras_cotas": IDENT + [
        ("grupo_padrao", "texto", "Grupo da cota, padronizado.", "pcd / lgbtqia / quilombolas / mulheres / povos_tradicionais / pessoas_idosas / territorio / outro", "modelo"),
        ("grupo", "texto", "Grupo da cota como escrito no edital.", "", "modelo"),
        ("percentual", "decimal", "Percentual reservado ao grupo.", "fração: 0.05 = 5%", "modelo"),
        ("vagas", "inteiro", "Vagas reservadas ao grupo, quando informadas.", "", "modelo"),
    ],
    "evidencias": IDENT + [
        ("campo", "texto", "Campo que o trecho sustenta.", "ex.: percentual_negras, vagas_totais, valor_total, objeto", "modelo"),
        ("trecho", "texto", "Trecho literal do edital (até ~300 caracteres).", "", "modelo"),
        ("pagina", "inteiro", "Página do PDF onde está o trecho.", "", "modelo"),
    ],
}

ABAS = {
    "editais": "Uma linha por PDF de edital. Dados extraídos, resultado da verificação da IN 5 e alertas de revisão.",
    "categorias": "Uma linha por categoria (módulo, linha ou faixa) de cada edital, com vagas, valores e reservas.",
    "outras_cotas": "Uma linha por reserva de vagas para grupos além de pessoas negras e indígenas.",
    "evidencias": "Uma linha por trecho do edital usado pelo modelo como evidência de um campo.",
}

SOBRE = [
    ("Arquivo", "data/final/lpg_editais.xlsx"),
    ("Conteúdo", "Editais de fomento da Lei Paulo Gustavo (LC nº 195/2022) publicados por estados, DF e capitais, "
                 "com a verificação das cotas da Instrução Normativa MinC nº 5/2023, art. 5º."),
    ("Regra verificada", "No mínimo 20% das vagas para pessoas negras e 10% para pessoas indígenas (caput); arredondamento "
                         "de frações ≥ 0,5 para cima e < 0,5 para baixo (§2º); cotas em todas as categorias (§3º); categoria "
                         "com uma única vaga pode ir para a ampla concorrência se o edital somar 20% e 10% (§4º)."),
    ("Extração", f"Modelo {settings.MODEL_ID}, prompt {settings.PROMPT_VERSION}, saída estruturada validada por schema. "
                 "PDFs com texto convertidos para Markdown (PyMuPDF); escaneados lidos como imagem."),
    ("Fonte 'metadado'", "Vem do arquivo ou da pasta do PDF."),
    ("Fonte 'modelo'", "Extraído pelo LLM do texto do edital. Vazio quando o edital não traz a informação."),
    ("Fonte 'calculado'", "Calculado em Python a partir do que o modelo extraiu (src/transform/conformidade.py e consolida_lpg.py)."),
    ("Fonte 'regex'", "Conferência por dicionário de termos (src/dicionario_cotas.py); serve só para apontar revisão manual."),
    ("Percentuais", "Sempre em fração decimal: 0.2 = 20%."),
    ("Valores", "Em reais (R$), sem formatação."),
    ("Ligação entre abas", "Pela coluna 'id' (e também por 'arquivo_pdf')."),
]


def gera_dicionario() -> None:
    planilha = pd.read_excel(settings.LPG_OUTPUT, sheet_name=None, nrows=0)
    for aba, df in planilha.items():
        documentadas = {c for c, *_ in DICIONARIO.get(aba, [])}
        faltando = [c for c in df.columns if c not in documentadas]
        sobrando = [c for c in documentadas if c not in df.columns]
        if faltando:
            print(f"⚠️ {aba}: colunas sem documentação: {faltando}")
        if sobrando:
            print(f"ℹ️ {aba}: documentadas mas ausentes na planilha: {sobrando}")

    abas = pd.DataFrame([{"aba": a, "descricao": d, "linhas": len(pd.read_excel(settings.LPG_OUTPUT, sheet_name=a, usecols=[0]))}
                         for a, d in ABAS.items()])
    colunas = pd.DataFrame([
        {"aba": aba, "coluna": c, "tipo": t, "descricao": d, "valores_ou_unidade": v, "fonte": f}
        for aba, campos in DICIONARIO.items() for c, t, d, v, f in campos
    ])
    with pd.ExcelWriter(SAIDA, engine="openpyxl") as xw:
        pd.DataFrame(SOBRE, columns=["item", "descricao"]).to_excel(xw, sheet_name="sobre", index=False)
        abas.to_excel(xw, sheet_name="abas", index=False)
        colunas.to_excel(xw, sheet_name="colunas", index=False)

    wb = load_workbook(SAIDA)
    larguras = {"sobre": {"A": 22, "B": 120}, "abas": {"A": 14, "B": 100, "C": 8},
                "colunas": {"A": 14, "B": 26, "C": 10, "D": 80, "E": 45, "F": 11}}
    for nome, cols in larguras.items():
        ws = wb[nome]
        for letra, w in cols.items():
            ws.column_dimensions[letra].width = w
        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        for cell in ws[1]:
            cell.font = Font(bold=True)
        ws.freeze_panes = "A2"
    wb["colunas"].auto_filter.ref = wb["colunas"].dimensions
    wb.save(SAIDA)
    print(f"📘 {SAIDA} ({len(colunas)} colunas documentadas)")
